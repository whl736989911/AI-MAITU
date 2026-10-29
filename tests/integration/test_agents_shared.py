"""Shared agent list and ownership behavior."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from tests.support.auth import create_user, resolve_user_id


@pytest.fixture
async def env(env_with_provider):
    yield env_with_provider


async def _shared_agent(env) -> tuple[object, dict[str, str], dict[str, str], str]:
    client, _srv, admin_auth = env
    owner_auth = await create_user(client, admin_auth, username="share_owner")
    peer_auth = await create_user(client, admin_auth, username="share_peer")
    response = await client.post(
        "/api/agents/from-expert/default",
        headers=owner_auth,
        json={"name": "shared-bot"},
    )
    assert response.status_code == 201, response.text
    return client, owner_auth, peer_auth, response.json()["agent_id"]


_OWNER_ONLY_CONTENT = {
    "memory.sqlite": "memory-db-secret",
    "checkpoints.sqlite": "checkpoints-db-secret",
    ".maitu/sessions/x.jsonl": "session-log-secret",
    ".maitu-auth/token.json": "auth-token-secret",
    ".env": "env-secret",
}

# An unmigrated workspace keeps the same files under the old directory names;
# the guard must protect those too.
_LEGACY_OWNER_ONLY_CONTENT = {
    ".octop/sessions/legacy.jsonl": "legacy-session-log-secret",
    ".octop-auth/token.json": "legacy-auth-token-secret",
}


async def _shared_agent_with_system_files(
    env,
) -> tuple[object, dict[str, str], dict[str, str], str]:
    """A shared agent whose workspace holds the owner-only system files."""
    client, owner_auth, peer_auth, agent_id = await _shared_agent(env)
    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": True},
    )
    assert response.status_code == 200, response.text
    for path, content in {**_OWNER_ONLY_CONTENT, **_LEGACY_OWNER_ONLY_CONTENT}.items():
        response = await client.put(
            f"/api/agents/{agent_id}/workspace/file",
            headers=owner_auth,
            params={"path": path},
            json={"content": content},
        )
        assert response.status_code == 200, response.text
    return client, owner_auth, peer_auth, agent_id


async def test_shared_agent_appears_for_other_user(env) -> None:
    client, owner_auth, peer_auth, agent_id = await _shared_agent(env)

    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": True},
    )
    assert response.status_code == 200, response.text
    assert response.json()["is_shared"] is True

    response = await client.get("/api/agents", headers=peer_auth)
    assert response.status_code == 200
    agents = {row["agent_id"]: row for row in response.json()}
    assert agent_id in agents
    assert agents[agent_id]["is_owner"] is False
    assert agents[agent_id]["owner_username"] == "share_owner"

    response = await client.get(f"/api/agents/{agent_id}", headers=peer_auth)
    assert response.status_code == 200
    assert response.json()["is_shared"] is True
    assert response.json()["is_owner"] is False


async def test_create_agent_can_be_shared(env) -> None:
    client, _srv, admin_auth = env
    owner_auth = await create_user(client, admin_auth, username="shared_create_owner")
    peer_auth = await create_user(client, admin_auth, username="shared_create_peer")

    response = await client.post(
        "/api/agents",
        headers=owner_auth,
        json={"name": "created-shared-bot", "config": {}, "is_shared": True},
    )
    assert response.status_code == 201, response.text
    agent_id = response.json()["agent_id"]
    assert response.json()["is_shared"] is True

    response = await client.get("/api/agents", headers=peer_auth)
    assert response.status_code == 200
    assert agent_id in {row["agent_id"] for row in response.json()}


async def test_peer_cannot_patch_shared_agent(env) -> None:
    client, owner_auth, peer_auth, agent_id = await _shared_agent(env)
    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": True},
    )
    assert response.status_code == 200, response.text

    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=peer_auth,
        json={"name": "hijack"},
    )
    assert response.status_code == 403


async def test_peer_cannot_mutate_shared_agent_workspace_or_skills(env) -> None:
    client, owner_auth, peer_auth, agent_id = await _shared_agent(env)
    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": True},
    )
    assert response.status_code == 200, response.text

    response = await client.put(
        f"/api/agents/{agent_id}/workspace/file",
        headers=peer_auth,
        params={"path": "SOUL.md"},
        json={"content": "peer mutation"},
    )
    assert response.status_code == 403

    response = await client.post(
        f"/api/agents/{agent_id}/skills/example/disable",
        headers=peer_auth,
    )
    assert response.status_code == 403

    response = await client.get(
        f"/api/agents/{agent_id}/skills",
        headers=peer_auth,
    )
    assert response.status_code == 200, response.text

    response = await client.post(
        "/api/mbti/apply",
        headers={**peer_auth, "X-Octop-Agent-Id": agent_id},
        json={"code": "INTJ"},
    )
    assert response.status_code == 403

    response = await client.get(
        f"/api/agents/{agent_id}/terminal/context",
        headers=peer_auth,
    )
    assert response.status_code == 403

    response = await client.get(
        f"/api/agents/{agent_id}/chat/welcome",
        headers=peer_auth,
    )
    assert response.status_code == 200, response.text
    assert "welcome_message" in response.json()

    response = await client.get(f"/api/agents/{agent_id}/threads", headers=peer_auth)
    assert response.status_code == 200, response.text


async def test_peer_can_read_shared_agent_inbound_attachment(env) -> None:
    client, owner_auth, peer_auth, agent_id = await _shared_agent(env)
    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": True},
    )
    assert response.status_code == 200, response.text

    response = await client.put(
        f"/api/agents/{agent_id}/workspace/file",
        headers=owner_auth,
        params={"path": "inbound/shared-attachment.txt"},
        json={"content": "shared attachment"},
    )
    assert response.status_code == 200, response.text

    response = await client.get(
        f"/api/agents/{agent_id}/workspace/file",
        headers=peer_auth,
        params={"path": "inbound/shared-attachment.txt"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["content"] == "shared attachment"

    response = await client.get(
        f"/api/agents/{agent_id}/workspace/download",
        headers=peer_auth,
        params={"path": "inbound/shared-attachment.txt"},
    )
    assert response.status_code == 200, response.text
    assert response.content == b"shared attachment"


async def test_unshare_hides_agent_from_peer(env) -> None:
    client, owner_auth, peer_auth, agent_id = await _shared_agent(env)
    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": True},
    )
    assert response.status_code == 200, response.text

    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": False},
    )
    assert response.status_code == 200, response.text
    assert response.json()["is_shared"] is False

    response = await client.get("/api/agents", headers=peer_auth)
    assert response.status_code == 200
    assert agent_id not in {row["agent_id"] for row in response.json()}

    response = await client.get(f"/api/agents/{agent_id}", headers=peer_auth)
    assert response.status_code == 403


async def test_shared_agent_threads_are_private(env) -> None:
    client, owner_auth, peer_auth, agent_id = await _shared_agent(env)
    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": True},
    )
    assert response.status_code == 200, response.text

    response = await client.post(f"/api/agents/{agent_id}/threads", headers=owner_auth)
    assert response.status_code == 201, response.text
    owner_thread_id = response.json()["thread_id"]

    response = await client.post(f"/api/agents/{agent_id}/threads", headers=peer_auth)
    assert response.status_code == 201, response.text
    peer_thread_id = response.json()["thread_id"]

    response = await client.get(f"/api/agents/{agent_id}/threads", headers=peer_auth)
    assert response.status_code == 200, response.text
    assert {thread["thread_id"] for thread in response.json()} == {peer_thread_id}

    response = await client.get(
        f"/api/agents/{agent_id}/threads/{owner_thread_id}/history",
        headers=peer_auth,
    )
    assert response.status_code == 403


async def test_peer_cannot_rebind_session_to_owner_thread(env) -> None:
    client, owner_auth, peer_auth, agent_id = await _shared_agent(env)
    response = await client.patch(
        f"/api/agents/{agent_id}",
        headers=owner_auth,
        json={"is_shared": True},
    )
    assert response.status_code == 200, response.text

    response = await client.post(f"/api/agents/{agent_id}/threads", headers=owner_auth)
    assert response.status_code == 201, response.text
    owner_thread_id = response.json()["thread_id"]

    response = await client.patch(
        f"/api/agents/{agent_id}/session",
        headers=peer_auth,
        json={"thread_id": owner_thread_id},
    )
    assert response.status_code == 403


async def test_a_directed_grant_puts_the_agent_in_the_list(env) -> None:
    """③: what the sharing API grants appears in the list the dashboard draws.

    The list was "owned ∪ published", so a directed grant, a unit share or a role
    share opened the agent on a direct ``GET /api/agents/{id}`` (200) and left it
    out of every list — and the dashboard only draws the list, which is why the
    grantee saw nothing at all. The rule is now the one that decides opening an
    agent (``sharing.allowed_resource_ids``), so the two cannot disagree.
    """
    client, _srv, admin_auth = env
    # The ACL endpoints are the user-administration surface (``users``), so the
    # granter has to hold it — the agent's own owner is who grants here.
    owner_auth = await create_user(
        client, admin_auth, username="grant_owner", permissions=["users", "experts"]
    )
    peer_auth = await create_user(client, admin_auth, username="grant_peer")
    stranger_auth = await create_user(client, admin_auth, username="grant_stranger")
    created = await client.post(
        "/api/agents/from-expert/default", headers=owner_auth, json={"name": "grantable-bot"}
    )
    assert created.status_code == 201, created.text
    agent_id = created.json()["agent_id"]
    peer_id = await resolve_user_id(client, admin_auth, "grant_peer")

    async def peer_list() -> set[str]:
        response = await client.get("/api/agents", headers=peer_auth)
        assert response.status_code == 200, response.text
        return {row["agent_id"] for row in response.json()}

    assert agent_id not in await peer_list()
    assert (await client.get(f"/api/agents/{agent_id}", headers=peer_auth)).status_code == 403

    granted = await client.post(
        f"/api/sharing/acl/agent/{agent_id}",
        headers=owner_auth,
        json={
            "visibility": "private",
            "grants": [{"grantee_type": "user", "grantee_id": str(peer_id)}],
        },
    )
    assert granted.status_code == 200, granted.text

    assert (await client.get(f"/api/agents/{agent_id}", headers=peer_auth)).status_code == 200
    assert agent_id in await peer_list()

    # The grant reaches that one account: nobody else's list grew.
    stranger_list = await client.get("/api/agents", headers=stranger_auth)
    assert stranger_list.status_code == 200, stranger_list.text
    assert agent_id not in {row["agent_id"] for row in stranger_list.json()}


async def test_shared_peer_downloads_only_its_agent_facing_artifact_namespace(
    env, monkeypatch: pytest.MonkeyPatch
) -> None:
    client, server, admin_auth = env
    owner_auth = await create_user(client, admin_auth, username="artifact_owner")
    peer_auth = await create_user(client, admin_auth, username="artifact_peer")
    created = await client.post(
        "/api/features",
        headers=owner_auth,
        json={"feature_id": "shared-artifact", "name": "shared-artifact-bot"},
    )
    assert created.status_code == 201, created.text
    agent_id = created.json()["agent_id"]
    document = Document()
    document.add_paragraph("共享功能生成的文档")
    output = BytesIO()
    document.save(output)
    body = output.getvalue()
    written = await client.post(
        f"/api/agents/{agent_id}/workspace/upload",
        headers=owner_auth,
        params={"path": "report.docx"},
        files={
            "file": (
                "report.docx",
                body,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert written.status_code == 200, written.text
    published = await client.put(
        f"/api/agents/{agent_id}/workflow",
        headers=owner_auth,
        json={
            "workflow": {
                "version": 1,
                "status": "active",
                "inputs": {
                    "type": "object",
                    "properties": {
                        "topic": {"type": "string", "title": {"zh": "主题", "en": "Topic"}}
                    },
                },
                "steps": [{"id": "generate", "name": "生成", "prompt": "生成报告"}],
                "outputs": [{"name": "报告", "form": "file"}],
            }
        },
    )
    assert published.status_code == 200, published.text
    peer_id = await resolve_user_id(client, admin_auth, "artifact_peer")
    shared = await client.post(
        "/api/sharing/acl/feature/shared-artifact",
        headers=admin_auth,
        json={
            "visibility": "private",
            "grants": [{"grantee_type": "user", "grantee_id": str(peer_id)}],
        },
    )
    assert shared.status_code == 200, shared.text

    registry = server.app_runtime.agent_registry
    original_get_config = registry.get_config
    facing_root = f"/.maitu/workspaces/{agent_id}"

    def get_config(requested_agent_id: str) -> dict[str, object]:
        if requested_agent_id == agent_id:
            return {"workspace_dir": facing_root}
        return original_get_config(requested_agent_id)

    monkeypatch.setattr(registry, "get_config", get_config)
    downloaded = await client.get(
        f"/api/agents/{agent_id}/workspace/download",
        headers=peer_auth,
        params={"path": f"{facing_root}/report.docx"},
    )
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.content == body
    assert Document(BytesIO(downloaded.content)).paragraphs[0].text == "共享功能生成的文档"

    foreign = await client.get(
        f"/api/agents/{agent_id}/workspace/download",
        headers=peer_auth,
        params={"path": "/.maitu/workspaces/another-agent/report.docx"},
    )
    assert foreign.status_code == 403, foreign.text

    protected = await client.get(
        f"/api/agents/{agent_id}/workspace/download",
        headers=peer_auth,
        params={"path": f"{facing_root}/.maitu/sessions/x.jsonl"},
    )
    assert protected.status_code == 403, protected.text


async def test_peer_cannot_read_shared_agent_system_files(env) -> None:
    client, _owner_auth, peer_auth, agent_id = await _shared_agent_with_system_files(env)

    for path in {**_OWNER_ONLY_CONTENT, **_LEGACY_OWNER_ONLY_CONTENT}:
        for route in ("file", "download"):
            response = await client.get(
                f"/api/agents/{agent_id}/workspace/{route}",
                headers=peer_auth,
                params={"path": path},
            )
            assert response.status_code == 403, (
                f"{route} {path}: {response.status_code} {response.text}"
            )

    # Windows ignores trailing dots/spaces and resolves NTFS stream syntax, so
    # aliases of the reserved names must be refused too.
    for path in (
        "memory.sqlite.",
        "checkpoints.sqlite.",
        ".env.",
        ".maitu.",
        ".octop.",
        "memory.sqlite::$DATA",
    ):
        response = await client.get(
            f"/api/agents/{agent_id}/workspace/file",
            headers=peer_auth,
            params={"path": path},
        )
        assert response.status_code == 403, f"alias {path}: {response.status_code} {response.text}"

    # The document and media-preview read routes enforce the same guard.
    for path in (".maitu/sessions/x.jsonl", ".octop/sessions/legacy.jsonl"):
        response = await client.get(
            f"/api/agents/{agent_id}/workspace/doc",
            headers=peer_auth,
            params={"path": path},
        )
        assert response.status_code == 403, f"doc {path}: {response.text}"
    response = await client.get(
        f"/api/agents/{agent_id}/media/preview",
        headers=peer_auth,
        params={"source": "memory.sqlite"},
    )
    assert response.status_code == 403, response.text


async def test_owner_can_read_shared_agent_system_files(env) -> None:
    client, owner_auth, _peer_auth, agent_id = await _shared_agent_with_system_files(env)

    for path, content in {**_OWNER_ONLY_CONTENT, **_LEGACY_OWNER_ONLY_CONTENT}.items():
        response = await client.get(
            f"/api/agents/{agent_id}/workspace/file",
            headers=owner_auth,
            params={"path": path},
        )
        assert response.status_code == 200, f"file {path}: {response.status_code} {response.text}"
        assert response.json()["content"] == content

        response = await client.get(
            f"/api/agents/{agent_id}/workspace/download",
            headers=owner_auth,
            params={"path": path},
        )
        assert response.status_code == 200, (
            f"download {path}: {response.status_code} {response.text}"
        )
        assert response.content == content.encode()


@pytest.mark.parametrize("storage", ["projection", "archive"])
async def test_shared_history_restores_virtual_and_inline_file_downloads(
    env, tmp_path, storage: str
) -> None:
    import base64
    import json

    from harness_agent.builtin.tools.send_file import build_send_file_to_user_tool
    from langchain_core.messages import HumanMessage, ToolMessage

    from octop.infra.gateway.process.history_projection import message_inputs
    from octop.infra.history.recorder import RecordingTracker
    from octop.infra.history.service import HistoryArchive
    from octop.infra.history.store import HistoryStore

    client, server, admin_auth = env
    _, owner_auth, peer_auth, agent_id = await _shared_agent(env)
    response = await client.patch(
        f"/api/agents/{agent_id}", headers=owner_auth, json={"is_shared": True}
    )
    assert response.status_code == 200, response.text
    workspace = server.app_runtime.agent_registry.get_agent(agent_id).workspace
    document = b"%PDF-generated-document"
    inline_data = b"%PDF-inline-document"
    await workspace.aupload_bytes("generated/report.pdf", document)
    tool = build_send_file_to_user_tool(workspace)
    block = await tool.ainvoke({"file_path": "/generated/report.pdf"})
    inline = {
        "type": "file",
        "filename": "inline.pdf",
        "source": {
            "type": "base64",
            "data": base64.b64encode(inline_data).decode("ascii"),
            "media_type": "application/pdf",
        },
    }
    messages = [
        HumanMessage(content="send both files", id="history-user"),
        ToolMessage(content=json.dumps(block), tool_call_id="send", name="send_file_to_user"),
        ToolMessage(content=json.dumps(inline), tool_call_id="inline", name="custom_file"),
    ]
    peer_id = await resolve_user_id(client, admin_auth, "share_peer")
    thread_id = f"file-history-{storage}"
    server.services.thread_repo.insert(
        thread_id=thread_id,
        agent_id=agent_id,
        user_id=peer_id,
        channel_type="dashboard",
        session_key=thread_id,
        last_active=0,
    )
    previous = server.app_runtime.history_archive
    store = None
    try:
        if storage == "archive":
            store = HistoryStore(tmp_path / "history.sqlite", identity="file-history")
            archive = HistoryArchive(
                store,
                server.services.thread_message_repo,
                server.services.trajectory_event_repo,
                enabled=True,
            )
            server.app_runtime.history_archive = archive
            turn = archive.begin(agent_id, thread_id)
            recorder = RecordingTracker(archive, turn, messages[:1])
            for message in messages[1:]:
                recorder.observe({"type": "tool_result", "messages": [message]})
            await recorder.finish(completed=True)
        else:
            server.app_runtime.history_archive = None
            server.services.thread_message_repo.replace_all(thread_id, message_inputs(messages))
        response = await client.get(
            f"/api/agents/{agent_id}/threads/{thread_id}/history", headers=peer_auth
        )
        assert response.status_code == 200, response.text
        files = {
            value["filename"]: value
            for message in response.json()["messages"]
            for item in message.get("content", [])
            if isinstance(item, dict) and item.get("type") == "tool_result"
            for value in [json.loads(item["output"])]
        }
        assert files["report.pdf"]["path"] == "generated/report.pdf"
        for filename, data in (("report.pdf", document), ("inline.pdf", inline_data)):
            downloaded = await client.get(
                f"/api/agents/{agent_id}/workspace/download",
                params={"path": files[filename]["path"]},
                headers=peer_auth,
            )
            assert downloaded.status_code == 200, downloaded.text
            assert downloaded.content == data
    finally:
        server.app_runtime.history_archive = previous
        if store is not None:
            store.close()
