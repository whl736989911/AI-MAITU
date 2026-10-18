"""Authorized, bounded reads of original knowledge files, independent of indexes."""

from __future__ import annotations

import asyncio
import base64
import hashlib
import json
import mimetypes
import threading
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from octop.infra.knowledge.files import file_digest
from octop.infra.knowledge.image_parts import image_part_count, read_image_parts
from octop.infra.knowledge.parse import ParsedBlock, PasswordRequiredError, parse_document
from octop.infra.knowledge.relpath import normalize_kb_path
from octop.infra.knowledge.service import KnowledgeService, knowledge_content_type
from octop.infra.knowledge.sources import is_folder_kind

_READ_LOCK = threading.RLock()
_MAX_TEXT = 40_000
_MAX_RANGE = 200
_MAX_IMAGES = 4
_MAX_IMAGE_BYTES = 8 * 1024 * 1024
_MAX_OFFICE_EXPANDED = 256 * 1024 * 1024
_TEXT_TYPES = frozenset({".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv"})
_IMAGE_TYPES = frozenset({".png", ".jpg", ".jpeg", ".webp"})


class KnowledgeSourceService:
    """Keep source access, evidence and workspace materialization under one ACL."""

    def __init__(self, services: Any) -> None:
        self._services = services
        self._knowledge = KnowledgeService(services)

    @staticmethod
    def supports_format(filename: str) -> bool:
        """Whether parsed source reads support the extension, not file integrity."""
        suffix = Path(filename).suffix.lower()
        return suffix in _TEXT_TYPES or knowledge_content_type(suffix) is not None

    def _document(self, kb_id: str, doc_id: str, actor_user_id: int, is_admin: bool) -> Any:
        self._knowledge.get_readable_base(kb_id, actor_user_id=actor_user_id, is_admin=is_admin)
        document = self._services.knowledge_repo.get_document(doc_id)
        if document is None or document.kb_id != kb_id or document.is_dir:
            raise LookupError("knowledge document not found")
        self._knowledge.require_document_readable(
            document, actor_user_id=actor_user_id, is_admin=is_admin
        )
        return document

    def describe_document(
        self, kb_id: str, doc_id: str, *, actor_user_id: int, is_admin: bool = False
    ) -> dict[str, str]:
        document = self._document(kb_id, doc_id, actor_user_id, is_admin)
        return {
            "document_id": document.id,
            "filename": document.filename,
            "content_type": document.content_type,
            "path": document.path,
        }

    def list_documents(
        self, kb_id: str, *, actor_user_id: int, is_admin: bool = False, path: str = ""
    ) -> dict[str, Any]:
        prefix = normalize_kb_path(path)
        if prefix:
            folder = self._services.knowledge_repo.get_document_by_path(kb_id, prefix)
            if folder is None or not folder.is_dir:
                raise LookupError("knowledge folder not found")
            self._knowledge.require_document_readable(
                folder, actor_user_id=actor_user_id, is_admin=is_admin
            )
        rows = self._knowledge.list_documents(
            kb_id, actor_user_id=actor_user_id, is_admin=is_admin, prefix=prefix or None
        )
        return {
            "kb_id": kb_id,
            "path": prefix,
            "documents": [
                {
                    "document_id": row.id,
                    "filename": row.filename,
                    "path": row.path,
                    "is_dir": row.is_dir,
                    "content_type": row.content_type,
                    "byte_size": row.byte_size,
                    "indexed": row.status == "ready",
                    "version": row.content_hash,
                }
                for row in rows
            ],
        }

    @contextmanager
    def open_original(
        self, kb_id: str, doc_id: str, *, actor_user_id: int, is_admin: bool = False
    ) -> Iterator[Path]:
        document = self._document(kb_id, doc_id, actor_user_id, is_admin)
        if document.data_source_id:
            from octop.infra.knowledge.data_sources import DataSourceService
            from octop.infra.knowledge.jobs import _source_path

            source = self._services.data_sources_repo.get(document.data_source_id)
            if source is None:
                raise FileNotFoundError("knowledge source unavailable")
            if is_folder_kind(source.kind):
                connector = DataSourceService(self._services).connector(source)
                with _source_path(connector, document.source_path, document.filename) as original:
                    self._document(kb_id, doc_id, actor_user_id, is_admin)
                    yield original
                return
        original, _, _ = self._knowledge.resolve_document_file(
            kb_id, doc_id, actor_user_id=actor_user_id, is_admin=is_admin
        )
        yield original

    @staticmethod
    def _evidence_id(doc_id: str, version: str, locator: dict[str, Any]) -> str:
        key = json.dumps([doc_id, version, locator], sort_keys=True, ensure_ascii=False)
        return "source:" + hashlib.sha256(key.encode()).hexdigest()

    @staticmethod
    def _office_size(path: Path) -> tuple[str, ...]:
        if path.suffix.lower() not in {".docx", ".xlsx", ".xlsm", ".pptx"}:
            return ()
        with ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > _MAX_OFFICE_EXPANDED:
                raise ValueError("knowledge document expanded size exceeds read limit")
            return tuple(
                info.filename
                for info in archive.infolist()
                if not info.is_dir()
                and info.filename.startswith(("word/media/", "ppt/media/", "xl/media/"))
            )

    def read_document(
        self,
        kb_id: str,
        doc_id: str,
        *,
        actor_user_id: int,
        is_admin: bool = False,
        start: int = 1,
        end: int | None = None,
        mode: str = "text",
        sheet: str | None = None,
    ) -> dict[str, Any]:
        if start < 1 or mode not in {"text", "image", "metadata"}:
            raise ValueError("invalid source read range or mode")
        stop = end if end is not None else start + (0 if mode == "image" else 49)
        if stop < start or stop - start + 1 > (_MAX_IMAGES if mode == "image" else _MAX_RANGE):
            raise ValueError("source read range exceeds limit")
        with (
            _READ_LOCK,
            self.open_original(
                kb_id, doc_id, actor_user_id=actor_user_id, is_admin=is_admin
            ) as original,
        ):
            document = self._document(kb_id, doc_id, actor_user_id, is_admin)
            version = file_digest(original)
            if mode == "metadata":
                self._document(kb_id, doc_id, actor_user_id, is_admin)
                return {
                    "status": "metadata",
                    "kb_id": kb_id,
                    "document_id": doc_id,
                    "filename": document.filename,
                    "path": document.path,
                    "version": version,
                    "byte_size": original.stat().st_size,
                    "text": "",
                    "locator": {"kind": "file"},
                    "evidence_id": "",
                    "blocks": [],
                    "images": [],
                    "verified": False,
                    "complete": False,
                    "collection_updates": {"updated": 0, "reason": "metadata_only"},
                }
            embedded_images = self._office_size(original)
            suffix = original.suffix.lower()
            images: list[dict[str, Any]] = []
            total: int | None
            if suffix in _IMAGE_TYPES and end is None:
                stop = start
            if mode == "image" or suffix in _IMAGE_TYPES:
                images, total = self._images(original, start, stop)
                stop = min(stop, total)
                blocks: list[ParsedBlock] = []
            elif suffix in _TEXT_TYPES:
                blocks, total = self._text(original, start, stop)
            elif suffix == ".pdf":
                blocks, total = self._pdf(original, start, stop)
            elif suffix in {".xlsx", ".xlsm"}:
                blocks, total = self._spreadsheet(original, start, stop, sheet)
            else:
                parsed = parse_document(original, source_path=document.path)
                available = parsed.blocks or (
                    ParsedBlock(parsed.text, "document", {"kind": "document"}),
                )
                blocks = list(available[start - 1 : stop])
                total = len(available)
            if file_digest(original) != version:
                raise ValueError("knowledge source changed while reading; read again")
            self._document(kb_id, doc_id, actor_user_id, is_admin)
            # A source read must not claim that existing indexed excerpts still describe
            # newly observed bytes. This does not enqueue or parse the rest of the file.
            if document.content_hash != version:
                self._services.knowledge_repo.update_document(
                    doc_id, content_hash=version, status="discovered", chunk_count=0
                )
            evidence_blocks: list[dict[str, Any]] = []
            budget = _MAX_TEXT
            truncated = False
            for block in blocks:
                if budget <= 0:
                    truncated = True
                    break
                text = block.text[:budget]
                locator = dict(block.locator)
                if len(text) != len(block.text):
                    locator["excerpt_characters"] = len(text)
                    truncated = True
                if block.heading:
                    locator["heading"] = block.heading
                evidence_blocks.append(
                    {
                        "text": text,
                        "locator": locator,
                        "evidence_id": self._evidence_id(doc_id, version, locator),
                    }
                )
                budget -= len(text)
            text = "\n\n".join(str(block["text"]) for block in evidence_blocks)
            locator = {
                "kind": "image" if images else ("page" if suffix == ".pdf" else "range"),
                "start": start,
                "end": stop,
                "positions": (
                    [image["locator"] for image in images]
                    if images
                    else [block["locator"] for block in evidence_blocks]
                ),
            }
            result: dict[str, Any] = {
                "status": "image_ready" if images else ("read" if text.strip() else "needs_visual"),
                "kb_id": kb_id,
                "document_id": doc_id,
                "filename": document.filename,
                "path": document.path,
                "version": version,
                "text": text,
                "locator": locator,
                "evidence_id": self._evidence_id(doc_id, version, locator),
                "blocks": evidence_blocks,
                "images": images,
                "embedded_images": len(embedded_images),
                "embedded_image_parts": embedded_images,
                "total": total,
                "truncated": truncated,
                "complete": bool(
                    total is not None and start == 1 and stop >= total and not truncated
                ),
                "next_start": stop + 1 if total is None or stop < total else None,
                "verified": bool(text.strip()),
            }
        from octop.infra.knowledge.collections import KnowledgeCollectionService

        result["collection_updates"] = KnowledgeCollectionService(
            self._services
        ).record_source_read(kb_id, doc_id, result, actor_user_id=actor_user_id, is_admin=is_admin)
        return result

    @staticmethod
    def _text(path: Path, start: int, stop: int) -> tuple[list[ParsedBlock], int | None]:
        def read_lines(
            encoding: str, errors: str = "strict"
        ) -> tuple[list[ParsedBlock], int | None]:
            blocks: list[ParsedBlock] = []
            number = 0
            with path.open(encoding=encoding, errors=errors) as stream:
                for number, line in enumerate(stream, 1):
                    if number > stop:
                        return blocks, None
                    if number >= start:
                        blocks.append(ParsedBlock(line.rstrip("\r\n"), "line", {"line": number}))
            return blocks, number

        try:
            return read_lines("utf-8-sig")
        except UnicodeDecodeError:
            return read_lines("gb18030", "replace")

    @staticmethod
    def _pdf(path: Path, start: int, stop: int) -> tuple[list[ParsedBlock], int]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise PasswordRequiredError("knowledge document requires a password")
        total = len(reader.pages)
        if start > total:
            raise ValueError("source page is outside the document")
        return [
            ParsedBlock(reader.pages[number - 1].extract_text() or "", "page", {"page": number})
            for number in range(start, min(stop, total) + 1)
        ], total

    @staticmethod
    def _spreadsheet(
        path: Path, start: int, stop: int, sheet: str | None
    ) -> tuple[list[ParsedBlock], int | None]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            current = workbook[sheet] if sheet else workbook.worksheets[0]
            total = current.max_row
            if total and start > total:
                raise ValueError("source row is outside the worksheet")
            blocks = [
                ParsedBlock(
                    "\t".join("" if value is None else str(value) for value in row),
                    "row",
                    {"sheet": current.title, "row": number},
                )
                for number, row in enumerate(
                    current.iter_rows(
                        min_row=start, max_row=min(stop, total) if total else stop, values_only=True
                    ),
                    start,
                )
            ]
            return blocks, total
        finally:
            workbook.close()

    @staticmethod
    def _images(path: Path, start: int, stop: int) -> tuple[list[dict[str, Any]], int]:
        suffix = path.suffix.lower()
        if suffix in _IMAGE_TYPES:
            return read_image_parts(
                path,
                byte_size=path.stat().st_size,
                max_bytes=_MAX_IMAGE_BYTES,
                mime_type=mimetypes.guess_type(path.name)[0] or "application/octet-stream",
                locator={"kind": "image"},
                start=start,
                stop=stop,
            )
        if suffix in {".docx", ".pptx", ".xlsx", ".xlsm"}:
            members = KnowledgeSourceService._office_size(path)
            images = []
            total = 0
            with ZipFile(path) as archive:
                for number, member in enumerate(members, 1):
                    size = archive.getinfo(member).file_size
                    if Path(member).suffix.lower() not in _IMAGE_TYPES:
                        count = 1
                    else:
                        with archive.open(member) as stream:
                            count = image_part_count(stream, size, _MAX_IMAGE_BYTES)
                    offset = total
                    total += count
                    if stop <= offset or start > total:
                        continue
                    if Path(member).suffix.lower() not in _IMAGE_TYPES:
                        raise ValueError(
                            "embedded image format requires processing the original file"
                        )
                    with archive.open(member) as stream:
                        selected, _ = read_image_parts(
                            stream,
                            byte_size=size,
                            max_bytes=_MAX_IMAGE_BYTES,
                            mime_type=mimetypes.guess_type(member)[0] or "application/octet-stream",
                            locator={"kind": "embedded_image", "image": number, "part": member},
                            start=max(1, start - offset),
                            stop=min(count, stop - offset),
                        )
                        images.extend(selected)
            if start > total:
                raise ValueError("source image is outside the document")
            return images, total
        if suffix != ".pdf":
            raise ValueError("image mode requires an image, PDF, or Office document")
        import pymupdf

        pdf_api: Any = pymupdf

        images = []
        with pdf_api.open(path) as document:
            if document.needs_pass:
                raise PasswordRequiredError("knowledge document requires a password")
            total = len(document)
            if start > total:
                raise ValueError("source page is outside the document")
            for number in range(start, min(stop, total) + 1):
                page = document[number - 1]
                scale = min(1.5, 1600 / max(page.rect.width, page.rect.height))
                data = page.get_pixmap(matrix=pdf_api.Matrix(scale, scale), alpha=False).tobytes(
                    "png"
                )
                if len(data) > _MAX_IMAGE_BYTES:
                    raise ValueError("rendered source page exceeds transfer limit")
                images.append(
                    {
                        "mime_type": "image/png",
                        "data_base64": base64.b64encode(data).decode("ascii"),
                        "locator": {"page": number},
                    }
                )
        return images, total

    async def export_document(
        self,
        kb_id: str,
        doc_id: str,
        *,
        workspace: Any,
        actor_user_id: int,
        expected_version: str,
        is_admin: bool = False,
    ) -> dict[str, Any]:
        def obtain() -> tuple[bytes, str, str, str]:
            with self.open_original(
                kb_id, doc_id, actor_user_id=actor_user_id, is_admin=is_admin
            ) as original:
                document = self._document(kb_id, doc_id, actor_user_id, is_admin)
                before = original.stat()
                data = original.read_bytes()
                after = original.stat()
                if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
                    raise ValueError("knowledge source changed while obtaining; try again")
                version = hashlib.sha256(data).hexdigest()
                if version != expected_version:
                    raise ValueError("knowledge source version changed; inspect metadata again")
                if document.content_hash != version:
                    self._services.knowledge_repo.update_document(
                        doc_id, content_hash=version, status="discovered", chunk_count=0
                    )
                suffix = Path(document.filename).suffix.lower()
                return data, version, document.filename, suffix

        data, version, filename, suffix = await asyncio.get_running_loop().run_in_executor(
            None, obtain
        )
        self._document(kb_id, doc_id, actor_user_id, is_admin)
        target = f"inbound/knowledge/{doc_id}/{version[:16]}{suffix}"
        await workspace.aupload_bytes(target, data)
        # Export is deliberately not a reading event and never updates collection facts.
        return {
            "status": "obtained",
            "kb_id": kb_id,
            "document_id": doc_id,
            "filename": filename,
            "version": version,
            "workspace_path": target,
            "byte_size": len(data),
            "read": False,
        }
